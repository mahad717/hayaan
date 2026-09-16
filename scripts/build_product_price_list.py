#!/usr/bin/env python3
"""Build editable Product | Price list (xlsx) for Hayaan from Sanguni.so scraped data."""
import sys, os

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (FONT_NAME, HEADER_BOLD, PRIMARY, NEUTRAL_600, NEUTRAL_900,
                  font_title, font_body, font_caption, fill_data_row,
                  setup_sheet, style_header_row, style_data_row)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = "/home/z/my-project/download/hayaan-product-price-list.xlsx"

# (name, price) — price None = not listed on Sanguni, user fills in
SECTIONS = [
    ("TIER 1 — Power, Charging & Audio (fastest rotation)", [
        ("Power Bank Green Lion Thunder 80,000mAh", 65.00),
        ("Power Bank Lion Ultra-Thin 20,000mAh 120W", 49.90),
        ("Power Bank Anker MagGo 10,000mAh", 48.00),
        ("Power Bank Anker Zolo 10,000mAh", 45.00),
        ("Power Bank Porodo Chrome Volt 10,000mAh", 25.00),
        ("Power Bank Porodo Pocket Volt 10,000mAh", 22.00),
        ("Power Bank Maxvolt 10,000mAh", 24.00),
        ("Power Bank Aston MagSafe 10,000mAh", 22.00),
        ("Power Bank Porodo 5,000mAh with Lightning", 15.00),
        ("Battery Case iPhone 16 Pro Max 7000mAh", 45.00),
        ("Battery Case iPhone 17 Pro Max 7000mAh", 56.00),
        ("Anker Prime 100W 3-Port GaN Wall Charger", 65.00),
        ("Apple 20W USB-C Power Adapter (Original)", 25.00),
        ("WiWU Helix 65W Fast Charger", 18.00),
        ("Adaptor USB-C 25W PD", 4.50),
        ("Cable Charge USB-C to Lightning 1M", 2.00),
        ("Cable Charge USB-C 1M", 3.00),
        ("Cable Charge 60W USB-C 1M", 22.00),
        ("Cable Charge 240W USB-C 2M", 25.00),
        ("Anker Boom 3i Outdoor Bluetooth Speaker", 16.70),
        ("Headphone EarPods USB-C Wired", 4.00),
        ("Earphone Samsung 3.5mm Jack", 1.50),
        ("Headphone Samsung Type-C", 18.00),
        ("Headphone iPhone Wired", 24.00),
        ("AirPods Pro 3 / Pro 2 / AirPods 4 ANC", None),
        ("Galaxy Buds 3 / 4 Pro / Buds FE", None),
        ("Maxwell Smart Watch (budget line)", 55.00),
        ("Digital Voltage Stabilizer Lightwave", 11.50),
    ]),
    ("TIER 2 — Smartphones & Wearables", [
        ("Samsung Galaxy A06", 110.00),
        ("Samsung Galaxy A07", 125.00),
        ("Samsung Galaxy A16", 118.00),
        ("Samsung Galaxy A17", 165.00),
        ("Samsung Galaxy A36 5G", 285.00),
        ("Samsung Galaxy A37 5G", 320.00),
        ("Samsung Galaxy S25 / S26 Ultra", None),
        ("iPhone 17 (Physical SIM)", 990.00),
        ("iPhone 16 Pro", 1120.00),
        ("iPhone 15 Pro Max", 1170.00),
        ("iPhone 16 Pro Max", 1240.00),
        ("iPhone Air", 1280.00),
        ("iPhone 17 Pro (Physical SIM)", 1420.00),
        ("iPhone 17 Pro Max (Physical SIM)", 1550.00),
        ("Amazfit Active 42mm Smart Watch", 130.00),
        ("Samsung Galaxy Watch 7 44mm", 190.00),
        ("Huawei Watch GT 6 GPS 46mm", 230.00),
        ("Apple Watch SE 40mm", 250.00),
        ("Apple Watch Series 9 45mm", 360.00),
        ("Apple Watch Series 10 46mm", 390.00),
        ("Apple Watch Series 11 46mm", 440.00),
    ]),
    ("TIER 3 — Computers, TV & Gaming (big ticket)", [
        ("HP EliteBook 840 G7 Notebook", 310.00),
        ("HP EliteBook 830 G7 Intel Core i7-10Gen", 340.00),
        ("HP EliteBook 830 G8 Intel Core i7-11Gen", 380.00),
        ("HP EliteBook 840 G9 Intel Core i7-12Gen", 950.00),
        ("Lenovo ThinkPad E14 / T14 / X13", None),
        ("Asus TUF FX507 Gaming i7-13Gen", 1320.00),
        ("MacBook Pro M3", 1350.00),
        ("MacBook Air M5", 1530.00),
        ("MacBook Pro M4", 1720.00),
        ("MacBook Pro M5", 1900.00),
        ("iPad 9 / 10 / 11 (Wi-Fi)", None),
        ("iPad Pro M4 / M5 11 / 13", None),
        ("Samsung Galaxy Tab A11", None),
        ("TV Geepas Smart 32", 175.00),
        ("TV Geepas Smart 43", 210.00),
        ("TV LG 43-Inch Smart", 450.00),
        ("TV LG 55-Inch Smart", 650.00),
        ("TV LG 65-Inch Smart", 850.00),
        ("TV LG 75-Inch Smart", 1290.00),
        ("TV Samsung 85-Inch Smart", 2350.00),
        ("PlayStation 5 Digital Edition (Slim)", 720.00),
        ("PlayStation Sony 5 Pro 2TB Slim", 1180.00),
        ("Xiaomi TV Stick / TV Box S 4K", None),
        ("Speaker Outdoor L208 100W", 70.00),
        ("Projector Maxwell Spectra", 120.00),
        ("Projector Epson CO-W01", 350.00),
        ("Projector Screen Manual 180x180cm", 90.00),
    ]),
    ("TIER 4 — Home, Office & Networking (steady sellers)", [
        ("Geepas Electric Kettle", 15.00),
        ("Geepas Blender 2in1", 25.00),
        ("Geepas Blender 4in1", 39.90),
        ("Sandwich Maker Geepas", 25.00),
        ("Popcorn Maker Geepas", 20.00),
        ("Rice Cooker Geepas with Steamer", 46.00),
        ("Geepas Coffee Maker 1.5L", 35.00),
        ("Geepas Mini Oven and Grill 48L", 75.00),
        ("Hot Air Styler TP-5+1", 10.00),
        ("Water Dispenser Light Wave", 120.00),
        ("Air Conditioner Lightwave 9K", 260.00),
        ("Air Conditioner Lightwave 12K", 280.00),
        ("Air Conditioner Lightwave 18K", 390.00),
        ("Air Conditioner Lightwave 24K", 460.00),
        ("TP-Link 840N Wi-Fi Router", 15.00),
        ("TP-Link AC750 Wireless Dual Band Router", 24.90),
        ("TP-Link AC1200 WiFi Range Extender", 36.00),
        ("TP-Link Omada AC1350 Access Point", 85.00),
        ("MikroTik L009 series Router", 135.00),
        ("UniFi AP AC PRO Wireless Access Point", 145.00),
        ("HP Printer Laser 107A", 75.00),
        ("HP Printer Laser MFP 135A", 165.00),
        ("Epson EcoTank L3210 / L3250", None),
        ("HP Toner Cartridge 651A Original Set", 277.00),
        ("Zebra Barcode Printer ZD220T", None),
        ("POS Receipt Printer TP600 80mm", None),
        ("Microphone DJI Mic Mini 2", 33.30),
        ("DJI Osmo Pocket 3 / 4 Creator Combo", None),
        ("GoPro Max 2 True 8K 360 Camera", None),
        ("Aputure Amaran 200XS Studio Light", 330.00),
        ("Battery for Canon LPE10 Normal", 18.00),
        ("Memory Card 128GB / 256GB Extreme Pro", None),
    ]),
]

wb = Workbook()
ws = wb.active
ws.title = "Product Price List"

headers = ["Product", "Price"]
last_col = len(headers) + 1  # B=2, C=3 -> last_col=3

setup_sheet(ws, title="Product Price List — In-Demand Electronics (Sanguni.so)", last_col=last_col)

# Header row (row 4)
for col_idx, h in enumerate(headers, start=2):
    ws.cell(row=4, column=col_idx, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

section_fill = PatternFill('solid', fgColor="D6E4F0")  # SECONDARY (primary light)
section_font = Font(name=FONT_NAME, size=11, bold=HEADER_BOLD, color=PRIMARY)

row = 5
data_row_index = 0
section_rows = []
for section_title, items in SECTIONS:
    # Section row (merged B:C)
    ws.cell(row=row, column=2, value=section_title)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
    for c in range(2, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 24
    section_rows.append(row)
    row += 1

    for name, price in items:
        ws.cell(row=row, column=2, value=name)
        pc = ws.cell(row=row, column=3)
        if price is not None:
            pc.value = price
        style_data_row(ws, row_num=row, col_start=2, col_end=last_col, row_index=data_row_index)
        pc.number_format = '$#,##0.00'
        pc.alignment = Alignment(horizontal='right', vertical='center')
        ws.row_dimensions[row].height = 22
        data_row_index += 1
        row += 1

last_data_row = row - 1

# Caption notes (Last+2, Last+3)
note_font = font_caption()
n1 = ws.cell(row=last_data_row + 2, column=2,
             value="Source: sanguni.so catalog (scraped 17 Sep 2026). Prices are Sanguni retail reference prices in USD.")
n1.font = note_font
n2 = ws.cell(row=last_data_row + 3, column=2,
             value="Blank price = not listed on Sanguni — fill in your own. Some prices converted from SOS at approx. 570 SOS/USD.")
n2.font = note_font

# Column widths
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 46
ws.column_dimensions['C'].width = 14

# Freeze header rows so scrolling keeps headers visible
ws.freeze_panes = 'A5'

wb.properties.creator = "Z.ai"
wb.save(OUT)
print(f"SAVED {OUT}")
print(f"Rows: title=1 header=1 sections={len(section_rows)} products={data_row_index}")
